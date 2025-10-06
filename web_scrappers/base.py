from abc import ABC, abstractmethod
import pandas as pd
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
import os


class WebScrapper(ABC):
    @abstractmethod
    def parse_url(url: str, **kargs) -> pd.DataFrame:
        pass

    def store_df_to_excel(
        self,
        df: pd.DataFrame,
        local_file_path: str,
        sheet_name: str,
        table_name: str,
    ) -> None:
        """
        Store a pandas Dataframe into an Excel file in a Table format.

        Args:
            df: pd.DataFrame -> DataFrame that will be stored in an excel file.
            local_file_path: str -> Path to the output Excel file. (e.g., 'path/to/excel.xlsx')
            sheet_name: str -> Name of the excel sheet where the data will be stored
            table_name: str -> Name of the excel table where the data will be stored

        Returns:
            None
        """
        string_params = [local_file_path, sheet_name, table_name]

        if not isinstance(df, pd.DataFrame):
            raise TypeError("df must be a pandas DataFrane.")
        if not all([isinstance(param, str) and param != "" for param in string_params]):
            raise TypeError(
                "local_file_path, sheet_name, and table_name parameters "
                "must be not null strings"
            )

        local_path = (
            "/".join(local_file_path.split("/")[:-1])
            if len(local_file_path.split("/")) > 1
            else local_file_path[:-5]  # takes out '.xlsx'
        )

        if not os.path.exists(local_path):
            raise ValueError(f"The directory {local_path} does not exist.")
        elif not local_file_path.endswith(".xlsx"):
            raise ValueError("The file name must end with '.xlsx'.")

        with pd.ExcelWriter(local_file_path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name=sheet_name)
            worksheet = writer.sheets[sheet_name]
            (max_row, max_col) = df.shape

            # Calculate the table range in Excel format (e.g., "A1:D10")
            table_ref = f"A1:{get_column_letter(max_col)}{max_row + 1}"

            table = Table(displayName=table_name, ref=table_ref)
            style = TableStyleInfo(
                name="TableStyleMedium9",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            table.tableStyleInfo = style
            worksheet.add_table(table)
